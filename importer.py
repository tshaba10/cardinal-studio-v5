"""Local-only CSV/XLSX -> a draft. No LLM calls, formulas or dataset generation."""

import csv
import io
import math
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from file_models import Requirements
from privacy import blocked_name, looks_sensitive

MAX_BYTES = 1_000_000
MAX_ROWS = 10_000
MAX_COLS = 50


def checked_workbook(data):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Choose a non-empty file smaller than 1 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 200 or sum(m.file_size for m in members) > 10_000_000:
                raise ValueError("Workbook expands beyond the safe POC size limit.")
            if any(
                any(
                    p in m.filename.lower() for p in ("vbaproject", "externallinks/", "embeddings/")
                )
                for m in members
            ):
                raise ValueError(
                    "Macros, external workbook links and embedded objects are not supported."
                )
        return load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not read this XLSX workbook. Save a plain .xlsx file and retry."
        ) from exc


def sheet_names(data):
    workbook = checked_workbook(data)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def normalize_header(value):
    name = re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")
    if not name:
        raise ValueError("Every column needs a non-empty header.")
    if not name[0].isalpha():
        name = "field_" + name
    if len(name) > 80:
        raise ValueError("Keep column names under 80 characters.")
    if blocked_name(name):
        raise ValueError(
            "A header is blocked by the POC safety policy. Remove card numbers, CVV/PIN, contact or credential columns; use tokens instead."
        )
    return name


def _excel_value(cell):
    if cell.data_type in ("f", "e"):
        raise ValueError(
            "Excel formulas and error cells are not supported. Upload plain values only."
        )
    value = cell.value
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time().isoformat() == "00:00:00"
            else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if type(value) is int and re.fullmatch(r"0{2,}", cell.number_format or ""):
        return str(value).zfill(len(cell.number_format))
    return value


def read_rows(data, filename, sheet=None, delimiter=","):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Choose a non-empty file smaller than 1 MB.")
    suffix = Path(filename).suffix.lower()
    workbook = None
    try:
        if suffix == ".csv":
            if delimiter not in (",", ";", "\t"):
                raise ValueError("Unsupported delimiter.")
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("CSV contains null bytes; save as UTF-8 CSV.")
            iterator = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        elif suffix == ".xlsx":
            workbook = checked_workbook(data)
            if sheet is None or sheet not in workbook.sheetnames:
                raise ValueError("Select a worksheet first.")
            worksheet = workbook[sheet]
            # Ignore incorrect exporter dimensions; bounds are enforced while reading.
            worksheet.reset_dimensions()
            iterator = ([_excel_value(cell) for cell in row] for row in worksheet.iter_rows())
        else:
            raise ValueError(
                "Use .csv or .xlsx. Old .xls and macro-enabled .xlsm are not supported."
            )
        header = next(iterator, [])
        if not header or len(header) > MAX_COLS:
            raise ValueError("The first row must contain 1–50 column headers.")
        if any(v is None or not str(v).strip() for v in header):
            raise ValueError("Every column needs a header in the first row.")
        names = [normalize_header(v) for v in header]
        if len(set(names)) != len(names):
            raise ValueError(
                "Duplicate headers (including names that normalize identically). Rename them first."
            )
        rows = []
        for number, row in enumerate(iterator, start=1):
            if number > MAX_ROWS:
                raise ValueError(
                    "This POC accepts at most 10,000 input rows; no silent truncation."
                )
            if not row or all(v is None or v == "" for v in row):
                continue
            if suffix == ".xlsx" and len(row) < len(names):
                row += [None] * (len(names) - len(row))
            if len(row) != len(names):
                raise ValueError(
                    f"Input row {number + 1} has a different number of cells than the headers."
                )
            if any(len(str(v)) > 2000 for v in row if v is not None):
                raise ValueError("Cells must be at most 2,000 characters.")
            if any(looks_sensitive(v) for v in row if v is not None):
                raise ValueError(
                    "Potentially sensitive values detected. Use de-identified/synthetic examples, not real cardholder data."
                )
            if any(isinstance(v, str) and v.lstrip().startswith(("=", "@")) for v in row):
                raise ValueError("Formula-like cells are not supported. Upload plain values only.")
            rows.append(row)
        return names, rows, [str(v) for v in header]
    except UnicodeError as exc:
        raise ValueError("Save the CSV as UTF-8 (CSV UTF-8 in Excel).") from exc
    except csv.Error as exc:
        raise ValueError("Malformed CSV. Check quoting and the selected delimiter.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not parse this file. Save plain UTF-8 CSV or a valid .xlsx workbook."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def infer_values(name, raw):
    if not raw:
        return "unknown", []
    strings = [str(v).strip() for v in raw]
    # Never lose leading zeroes in identifiers, postal codes or labels.
    if name.endswith("_id") or "postal" in name or "zip" in name:
        return "string", strings
    if all(type(v) is bool for v in raw) or all(v.lower() in ("true", "false") for v in strings):
        return "boolean", [v.lower() == "true" for v in strings]
    if all(re.fullmatch(r"-?(0|[1-9]\d*)", v) for v in strings):
        numbers = [int(v) for v in strings]
        if all(-32768 <= v <= 32767 for v in numbers):
            return "int16", numbers
        # Large integers might be identifiers; do not silently lose precision.
        return "string", strings
    if all(re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", v) for v in strings):
        numbers = [float(v) for v in strings]
        if all(math.isfinite(v) for v in numbers):
            return "decimal", numbers
    try:
        if all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in strings):
            return "date", [date.fromisoformat(v).isoformat() for v in strings]
        if all("T" in v and datetime.fromisoformat(v).tzinfo is not None for v in strings):
            return "timestamp", [datetime.fromisoformat(v).isoformat() for v in strings]
    except ValueError:
        pass
    return "string", strings


def import_draft(data, filename, sheet=None, delimiter=","):
    names, rows, original = read_rows(data, filename, sheet, delimiter)
    draft = Requirements().model_dump()
    report = []
    for index, name in enumerate(names):
        raw = [row[index] for row in rows if row[index] is not None and row[index] != ""]
        kind, values = infer_values(name, raw)
        counts = Counter(values)
        null_rate = (len(rows) - len(raw)) / len(rows) if rows else 0.0
        too_many = len(counts) > 100
        proposed_values = [] if too_many else list(counts)
        weights = [] if too_many else [count / len(raw) for count in counts.values()]
        draft["elements"].append(
            dict(
                col_name=name,
                type=kind,
                nullable=null_rate > 0,
                null_rate=null_rate,
                unique=False,
                values=proposed_values,
                weights=weights,
                provenance="Uploaded sample only; observed frequencies are proposals, not industry estimates.",
            )
        )
        report.append(
            dict(
                original_header=original[index],
                config_name=name,
                inferred_type=kind,
                nonempty_samples=len(raw),
                distinct_values=len(counts),
                missing_samples=len(rows) - len(raw),
                attention="Over 100 values: choose a smaller explicit domain or wait for a richer generator contract."
                if too_many
                else "Review type, allowed values, probabilities, missingness and uniqueness.",
            )
        )
    draft["assumptions"] = [
        "Uploaded rows are a small sample, not a representative industry population.",
        "Proposed frequencies and missingness describe the sample only; review before confirming.",
        "Fields are independently sampled from finite allowed values. Row relationships are NOT inferred.",
        "Uniqueness is NOT inferred from a few distinct sample values.",
        "Blank columns remain unresolved. No values or output row count have been invented.",
    ]
    return Requirements.model_validate(draft).model_dump(), report
